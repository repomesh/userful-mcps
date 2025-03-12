"""
# /// script
# dependencies = [
#     "python-docx",
#     "python-docx-replace",
#     "mcp",
#     "docx2pdf",
#     "docx>=0.2.4",
# ]
# requires-python = ">=3.10"
# ///
"""

import os
import base64
import json
import tempfile
import logging
from typing import Dict, Optional, Union, Any
import uuid

from docx import Document
from docx2pdf import convert
from python_docx_replace import docx_replace, docx_blocks, docx_get_keys
from mcp.server import Server
from mcp.server.stdio import stdio_server
from mcp.types import TextContent, Tool
from pydantic import BaseModel


class ProcessTemplate(BaseModel):
    template_file: str
    replacements: Dict[str, str]
    blocks: Optional[Dict[str, bool]] = None
    output_filename: Optional[str] = None


class GetTemplateKeys(BaseModel):
    template_file: str


class ConvertToPdf(BaseModel):
    docx_file: str
    pdf_output: Optional[str] = None


class DocxTools(str):
    PROCESS_TEMPLATE = "process_template"
    GET_TEMPLATE_KEYS = "get_template_keys"
    CONVERT_TO_PDF = "convert_to_pdf"


async def serve() -> None:
    logger = logging.getLogger(__name__)
    server = Server("mcp-docx-template")

    @server.list_tools()
    async def list_tools() -> list[Tool]:
        return [
            Tool(
                name=DocxTools.PROCESS_TEMPLATE,
                description="Process a Word document template by replacing placeholders and managing blocks",
                inputSchema=ProcessTemplate.model_json_schema(),
            ),
            Tool(
                name=DocxTools.GET_TEMPLATE_KEYS,
                description="Extract all replacement keys from a Word document template",
                inputSchema=GetTemplateKeys.model_json_schema(),
            ),
            Tool(
                name=DocxTools.CONVERT_TO_PDF,
                description="Convert a Word document (docx) to PDF format",
                inputSchema=ConvertToPdf.model_json_schema(),
            ),
        ]

    def _is_base64(data: str) -> bool:
        """Check if the input string is base64 encoded."""
        try:
            return base64.b64encode(base64.b64decode(data)).decode() == data
        except Exception:
            return False

    def _handle_input_file(file_data: str) -> tuple[str, bool]:
        """Handle input file which can be a path or base64 content.
        Returns the file path and a flag indicating if it's a temp file.
        """
        if os.path.exists(file_data):
            return file_data, False

        if _is_base64(file_data):
            # Create a temporary file for the base64 content
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            temp_file.write(base64.b64decode(file_data))
            temp_file.close()
            return temp_file.name, True

        raise ValueError("Input must be a valid file path or base64-encoded content")

    def _handle_output_file(output_filename: Optional[str], suffix: str) -> str:
        """Create an output file path based on the provided filename or generate a random one."""
        if output_filename:
            return output_filename
        return f"{uuid.uuid4()}{suffix}"

    def _encode_file_if_needed(
        file_path: str, return_base64: bool
    ) -> Union[str, Dict[str, Any]]:
        """Encode the file as base64 if needed or return the file path."""
        if return_base64:
            with open(file_path, "rb") as f:
                content = base64.b64encode(f.read()).decode()
            return {"base64_content": content}
        return {"file_path": file_path}

    @server.call_tool()
    async def call_tool(name: str, arguments: dict) -> list[TextContent]:
        try:
            if name == DocxTools.PROCESS_TEMPLATE:
                template_file = arguments["template_file"]
                replacements = arguments["replacements"]
                blocks = arguments.get("blocks")
                output_filename = arguments.get("output_filename")

                file_path, is_temp = _handle_input_file(template_file)

                # Process the template
                doc = Document(file_path)
                docx_replace(doc, **replacements)

                if blocks:
                    docx_blocks(doc, **blocks)

                # Save the processed document
                output_path = _handle_output_file(output_filename, ".docx")
                doc.save(output_path)

                # Clean up temporary file if needed
                if is_temp:
                    os.unlink(file_path)

                # Determine if we should return base64 or file path
                return_base64 = not os.path.exists(template_file)
                result = _encode_file_if_needed(output_path, return_base64)

                if return_base64:
                    os.unlink(
                        output_path
                    )  # Remove temp output file if returning base64

                return [TextContent(type="text", text=json.dumps(result))]

            elif name == DocxTools.GET_TEMPLATE_KEYS:
                template_file = arguments["template_file"]
                file_path, is_temp = _handle_input_file(template_file)

                # Extract keys from the template
                doc = Document(file_path)
                keys = docx_get_keys(doc)

                # Clean up temporary file if needed
                if is_temp:
                    os.unlink(file_path)

                return [TextContent(type="text", text=json.dumps({"keys": list(keys)}))]

            elif name == DocxTools.CONVERT_TO_PDF:
                docx_file = arguments["docx_file"]
                pdf_output = arguments.get("pdf_output")

                file_path, is_temp = _handle_input_file(docx_file)

                # Determine output PDF path
                if pdf_output:
                    output_path = pdf_output
                else:
                    output_path = os.path.splitext(file_path)[0] + ".pdf"

                # Convert DOCX to PDF
                convert(file_path, output_path)

                # Clean up temporary file if needed
                if is_temp:
                    os.unlink(file_path)

                # Determine if we should return base64 or file path
                return_base64 = not os.path.exists(docx_file)
                result = _encode_file_if_needed(output_path, return_base64)

                if return_base64:
                    os.unlink(
                        output_path
                    )  # Remove temp output file if returning base64

                return [TextContent(type="text", text=json.dumps(result))]

            else:
                raise ValueError(f"Unknown tool: {name}")

        except Exception as e:
            return [TextContent(type="text", text=json.dumps({"error": str(e)}))]

    options = server.create_initialization_options()
    async with stdio_server() as (read_stream, write_stream):
        await server.run(read_stream, write_stream, options, raise_exceptions=True)


if __name__ == "__main__":
    import asyncio

    asyncio.run(serve())
